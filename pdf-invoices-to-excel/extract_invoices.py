"""Read a folder of PDF invoices and write them to one Excel workbook.

Usage:
    python extract_invoices.py [INPUT_FOLDER] [OUTPUT_XLSX]

Defaults to sample_invoices/ and output/invoices.xlsx next to this script.
"""

import argparse
import logging
import re
from datetime import datetime
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
log = logging.getLogger("extract_invoices")

MONEY = r"-?\$?\d[\d,]*\.\d{2}\b"
DATE_FORMATS = ["%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%B %d, %Y", "%b %d, %Y", "%d %B %Y"]
DATE_PATTERN = (r"\d{1,2}/\d{1,2}/\d{2,4}"
                r"|\d{4}-\d{2}-\d{2}"
                r"|[A-Z][a-z]+\.? \d{1,2}, \d{4}"
                r"|\d{1,2} [A-Z][a-z]+ \d{4}")

INVOICE_NO_RE = re.compile(
    r"\b(?:invoice|inv\.?)\s*(?:no\.?|number|#)\s*:?\s*#?\s*([A-Z0-9][A-Z0-9-]*)", re.I)
INVOICE_DATE_RE = re.compile(rf"(?<!due )(?:invoice )?date\s*:\s*({DATE_PATTERN})", re.I)
SUBTOTAL_RE = re.compile(r"^sub-?\s?total\b", re.I)
TAX_RE = re.compile(r"^(?:sales\s+)?tax\b", re.I)
TOTAL_RE = re.compile(r"^(?:total due|total|amount due|balance due)\b", re.I)

# One line item: description, qty, unit price, amount (or qty first).
NUM = r"-?\d+(?:\.\d+)?"
ITEM_QTY_LAST = re.compile(rf"^(?P<desc>.+?)\s+(?P<qty>{NUM})\s+(?P<price>{MONEY})\s+(?P<amount>{MONEY})$")
ITEM_QTY_FIRST = re.compile(rf"^(?P<qty>{NUM})\s+(?P<desc>.+?)\s+(?P<price>{MONEY})\s+(?P<amount>{MONEY})$")


def to_number(text):
    return float(text.replace("$", "").replace(",", ""))


def parse_date(text):
    text = text.replace(".", "")
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def last_amount(line):
    found = re.findall(MONEY, line)
    return to_number(found[-1]) if found else None


def parse_line_items(lines):
    items = []
    in_table = False
    qty_first = False
    for line in lines:
        lower = line.lower()
        if not in_table:
            if "qty" in lower and ("amount" in lower or "total" in lower):
                in_table = True
                qty_first = lower.startswith("qty")
            continue
        if SUBTOTAL_RE.match(line):
            break
        match = (ITEM_QTY_FIRST if qty_first else ITEM_QTY_LAST).match(line)
        if match:
            items.append({
                "description": match["desc"].strip(),
                "qty": float(match["qty"]),
                "unit_price": to_number(match["price"]),
                "amount": to_number(match["amount"]),
            })
    return items


def parse_invoice(text):
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    vendor = re.sub(r"\binvoice\b", "", lines[0], flags=re.I).strip() if lines else ""

    inv = {"vendor": vendor, "invoice_no": None, "date": None,
           "subtotal": None, "tax": None, "total": None}

    if m := INVOICE_NO_RE.search(text):
        inv["invoice_no"] = m.group(1)
    if m := INVOICE_DATE_RE.search(text):
        inv["date"] = parse_date(m.group(1))

    for line in lines:
        if inv["subtotal"] is None and SUBTOTAL_RE.match(line):
            inv["subtotal"] = last_amount(line)
        elif inv["tax"] is None and TAX_RE.match(line):
            inv["tax"] = last_amount(line)
        elif inv["total"] is None and TOTAL_RE.match(line):
            inv["total"] = last_amount(line)

    inv["items"] = parse_line_items(lines)
    return inv


def check_invoice(inv):
    """Return a list of problems; an empty list means the numbers add up."""
    problems = []
    for field in ("invoice_no", "date", "total"):
        if inv[field] is None:
            problems.append(f"missing {field.replace('_', ' ')}")
    if not inv["items"]:
        problems.append("no line items found")
    if inv["subtotal"] is not None and inv["items"]:
        lines_sum = round(sum(i["amount"] for i in inv["items"]), 2)
        if abs(lines_sum - inv["subtotal"]) > 0.01:
            problems.append(f"line items sum to {lines_sum:.2f}, subtotal says {inv['subtotal']:.2f}")
    if None not in (inv["subtotal"], inv["tax"], inv["total"]):
        if abs(inv["subtotal"] + inv["tax"] - inv["total"]) > 0.01:
            problems.append("subtotal + tax does not equal total")
    return problems


def read_pdf_text(path):
    with pdfplumber.open(path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


# --- Excel output ----------------------------------------------------------

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
CURRENCY = '#,##0.00'


def write_sheet(ws, headers, rows, widths, currency_cols=(), date_cols=()):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for col_idx, width in enumerate(widths, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = width
        for cell in ws[letter][1:]:
            if col_idx in currency_cols:
                cell.number_format = CURRENCY
            elif col_idx in date_cols:
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def write_workbook(invoices, failures, out_path):
    wb = Workbook()

    ws = wb.active
    ws.title = "Invoices"
    rows = [[inv["file"], inv["vendor"], inv["invoice_no"], inv["date"],
             inv["subtotal"], inv["tax"], inv["total"], len(inv["items"]),
             "; ".join(inv["problems"]) or "OK"] for inv in invoices]
    write_sheet(ws, ["File", "Vendor", "Invoice No", "Invoice Date", "Subtotal",
                     "Tax", "Total", "Lines", "Check"],
                rows, [30, 30, 15, 13, 12, 10, 12, 7, 28],
                currency_cols=(5, 6, 7), date_cols=(4,))

    ws = wb.create_sheet("Line Items")
    rows = [[inv["vendor"], inv["invoice_no"], inv["date"], item["description"],
             item["qty"], item["unit_price"], item["amount"]]
            for inv in invoices for item in inv["items"]]
    write_sheet(ws, ["Vendor", "Invoice No", "Invoice Date", "Description",
                     "Qty", "Unit Price", "Amount"],
                rows, [30, 15, 13, 44, 7, 12, 12],
                currency_cols=(6, 7), date_cols=(3,))

    ws = wb.create_sheet("Problems")
    write_sheet(ws, ["File", "Problem"], failures, [32, 70])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("input_folder", nargs="?", default=HERE / "sample_invoices", type=Path)
    parser.add_argument("output_xlsx", nargs="?", default=HERE / "output" / "invoices.xlsx", type=Path)
    args = parser.parse_args()

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    log_path = args.output_xlsx.with_name("extract_log.txt")
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s  %(levelname)-7s %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
        handlers=[logging.FileHandler(log_path, mode="w", encoding="utf-8"),
                  logging.StreamHandler()],
    )

    pdf_files = sorted(args.input_folder.glob("*.pdf"), key=lambda p: p.name.lower())
    if not pdf_files:
        log.error("No PDF files found in %s", args.input_folder)
        return

    invoices, failures = [], []
    for path in pdf_files:
        try:
            text = read_pdf_text(path)
        except Exception as exc:  # damaged or encrypted files
            log.warning("%s: could not open (%s)", path.name, exc)
            failures.append([path.name, f"could not open: {exc}"])
            continue

        if len(text.strip()) < 20:
            msg = "no text layer, probably a scan. Needs OCR or manual entry."
            log.warning("%s: %s", path.name, msg)
            failures.append([path.name, msg])
            continue

        inv = parse_invoice(text)
        inv["file"] = path.name
        inv["problems"] = check_invoice(inv)
        invoices.append(inv)

        if inv["problems"]:
            msg = "; ".join(inv["problems"])
            log.warning("%s: parsed with problems: %s", path.name, msg)
            failures.append([path.name, msg])
        else:
            log.info("%s: %s %s, total %.2f, %d lines", path.name, inv["vendor"],
                     inv["invoice_no"], inv["total"], len(inv["items"]))

    invoices.sort(key=lambda i: (i["vendor"].lower(), i["date"] or datetime.max.date()))
    write_workbook(invoices, failures, args.output_xlsx)
    log.info("Wrote %d invoices to %s. Files needing attention: %d",
             len(invoices), args.output_xlsx.name, len(failures))


if __name__ == "__main__":
    main()
