"""Clean up a messy contacts spreadsheet and remove duplicates.

Usage:
    python clean_contacts.py messy_contacts.csv
    python clean_contacts.py contacts.xlsx --out output/contacts_clean.xlsx

Expects these columns (any order, header names are matched loosely):
Name, Company, Email, Phone, City, State, Last Contact
"""

import argparse
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
COLUMNS = ["Name", "Company", "Email", "Phone", "City", "State", "Last Contact"]

EMAIL_RE = re.compile(r"^[a-z0-9._%+-]+@[a-z0-9-]+(\.[a-z0-9-]+)*\.[a-z]{2,}$")
DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%m-%d-%Y", "%b %d, %Y", "%B %d, %Y", "%m/%d/%y"]

US_STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA",
    "colorado": "CO", "connecticut": "CT", "delaware": "DE", "district of columbia": "DC",
    "florida": "FL", "georgia": "GA", "hawaii": "HI", "idaho": "ID", "illinois": "IL",
    "indiana": "IN", "iowa": "IA", "kansas": "KS", "kentucky": "KY", "louisiana": "LA",
    "maine": "ME", "maryland": "MD", "massachusetts": "MA", "michigan": "MI", "minnesota": "MN",
    "mississippi": "MS", "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV",
    "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM", "new york": "NY",
    "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK", "oregon": "OR",
    "pennsylvania": "PA", "rhode island": "RI", "south carolina": "SC", "south dakota": "SD",
    "tennessee": "TN", "texas": "TX", "utah": "UT", "vermont": "VT", "virginia": "VA",
    "washington": "WA", "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY",
}


# --- Field cleaners ---------------------------------------------------------

def squeeze(value):
    """Trim and collapse repeated spaces."""
    return re.sub(r"\s+", " ", str(value)).strip()


def fix_case(text):
    """Title-case text that was typed in all lower or all upper case.

    Mixed-case input is left alone so names like "MacLeod" or "DeShawn"
    are not damaged.
    """
    if not text or not (text.islower() or text.isupper()):
        return text

    def cap(word):
        word = word.capitalize()
        if word.startswith("Mc") and len(word) > 2:
            word = "Mc" + word[2:].capitalize()
        return word

    # Capitalize after spaces, hyphens and apostrophes: o'connell -> O'Connell
    return re.sub(r"[^\s'-]+", lambda m: cap(m.group()), text.lower())


def clean_email(value):
    email = squeeze(value).lower().replace("mailto:", "")
    if not email:
        return "", None
    if not EMAIL_RE.match(email):
        return email, "invalid email"
    return email, None


def clean_phone(value):
    raw = squeeze(value)
    if not raw:
        return "", None
    ext = ""
    m = re.search(r"\s*(?:x|ext\.?|extension)\s*(\d+)$", raw, re.I)
    if m:
        ext = f" x{m.group(1)}"
        raw = raw[:m.start()]
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) != 10:
        return squeeze(value), f"phone has {len(digits)} digits"
    return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}{ext}", None


def clean_state(value):
    state = squeeze(value)
    if len(state) == 2:
        return state.upper()
    return US_STATES.get(state.lower(), state)


def clean_date(value):
    text = squeeze(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def clean_row(raw):
    email, email_issue = clean_email(raw["Email"])
    phone, phone_issue = clean_phone(raw["Phone"])
    last_contact = clean_date(raw["Last Contact"])
    issues = [i for i in (email_issue, phone_issue) if i]
    if raw["Last Contact"].strip() and last_contact is None:
        issues.append("unreadable date")
    return {
        "Name": fix_case(squeeze(raw["Name"])),
        "Company": squeeze(raw["Company"]),
        "Email": email,
        "Phone": phone,
        "City": fix_case(squeeze(raw["City"])),
        "State": clean_state(raw["State"]),
        "Last Contact": last_contact,
        "issues": issues,
    }


# --- Duplicate handling -----------------------------------------------------

def phone_digits(phone):
    return re.sub(r"\D", "", phone)[:10]


def find_duplicates(rows):
    """Return (kept, removed). Each removed entry says why it went.

    A row is a duplicate when its email matches an earlier row (ignoring case
    and spaces), or when it has no matching email but the same name and phone.
    The first row is kept; blanks in it are filled from the duplicate, and
    the newer Last Contact date wins.
    """
    kept, removed = [], []
    by_email, by_name_phone = {}, {}

    for row in rows:
        src = row["source_row"]
        if not any(row["raw"][c].strip() for c in COLUMNS):
            removed.append({**row, "reason": "Empty row", "kept_row": None})
            continue

        email_key = row["Email"]
        name_phone_key = (row["Name"].lower(), phone_digits(row["Phone"])) if row["Phone"] else None
        match = by_email.get(email_key) if email_key else None
        how = "email"
        if match is None and name_phone_key:
            match = by_name_phone.get(name_phone_key)
            how = "name+phone"

        if match is None:
            kept.append(row)
            if email_key:
                by_email[email_key] = row
            if name_phone_key:
                by_name_phone[name_phone_key] = row
            continue

        same = all(row[c] == match[c] for c in COLUMNS)
        if same:
            reason = f"Exact duplicate of row {match['source_row']} after cleanup"
        elif how == "email":
            reason = f"Same email as row {match['source_row']}"
            a, b = row["raw"]["Email"], match["raw"]["Email"]
            if a.strip() != b.strip() and a.strip().lower() == b.strip().lower():
                reason += " (different capitalization)"
            elif a != b:
                reason += " (extra spaces)"
        else:
            reason = f"Same name and phone as row {match['source_row']}"

        merged = []
        for col in COLUMNS:
            if not match[col] and row[col]:
                match[col] = row[col]
                merged.append(col)
        if row["Last Contact"] and match["Last Contact"] and row["Last Contact"] > match["Last Contact"]:
            match["Last Contact"] = row["Last Contact"]
            merged.append("Last Contact")
        if merged:
            reason += f". Copied {', '.join(merged)} to the kept row"

        removed.append({**row, "reason": reason, "kept_row": match["source_row"]})
        if name_phone_key:
            by_name_phone.setdefault(name_phone_key, match)

    return kept, removed


# --- Input / output ---------------------------------------------------------

def load_rows(path):
    if path.suffix.lower() in (".xlsx", ".xls"):
        df = pd.read_excel(path, dtype=str, keep_default_na=False)
    else:
        df = pd.read_csv(path, dtype=str, keep_default_na=False, skip_blank_lines=False)

    # Match headers loosely: " email address" -> Email, "last_contact" -> Last Contact
    lookup = {re.sub(r"[^a-z]", "", c.lower()): c for c in df.columns}
    aliases = {"emailaddress": "email", "phonenumber": "phone", "fullname": "name"}
    for key, target in aliases.items():
        if key in lookup and target not in lookup:
            lookup[target] = lookup[key]
    missing = [c for c in COLUMNS if re.sub(r"[^a-z]", "", c.lower()) not in lookup]
    if missing:
        raise SystemExit(f"Missing columns: {', '.join(missing)}")

    rows = []
    for i, record in enumerate(df.to_dict("records")):
        raw = {c: record[lookup[re.sub(r'[^a-z]', '', c.lower())]] for c in COLUMNS}
        rows.append({**clean_row(raw), "raw": raw, "source_row": i + 2})  # +2: header is row 1
    return rows


HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
FLAG_FILL = PatternFill("solid", fgColor="FFF2B3")


def add_sheet(wb, title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def write_output(kept, removed, out_path):
    wb = Workbook()
    ws = add_sheet(wb, "Contacts", COLUMNS + ["Source Row", "Issues"],
                   [20, 26, 36, 20, 15, 7, 13, 11, 26], first=True)
    for row in sorted(kept, key=lambda r: r["Name"].split()[-1].lower() if r["Name"] else ""):
        ws.append([row[c] for c in COLUMNS] + [row["source_row"], "; ".join(row["issues"])])
        r = ws.max_row
        ws.cell(r, 7).number_format = "yyyy-mm-dd"
        for issue in row["issues"]:
            col = 3 if "email" in issue else 4 if "phone" in issue else 7
            ws.cell(r, col).fill = FLAG_FILL
    ws.auto_filter.ref = ws.dimensions

    ws = add_sheet(wb, "Removed Rows", ["Source Row", "Name (as entered)", "Email (as entered)",
                                        "Reason", "Kept Row"], [11, 22, 36, 72, 10])
    for row in removed:
        ws.append([row["source_row"], row["raw"]["Name"], row["raw"]["Email"],
                   row["reason"], row["kept_row"]])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description="Clean and de-duplicate a contacts sheet.")
    parser.add_argument("input", type=Path, help="CSV or Excel file")
    parser.add_argument("--out", type=Path, default=HERE / "output" / "contacts_clean.xlsx")
    args = parser.parse_args()

    rows = load_rows(args.input)
    kept, removed = find_duplicates(rows)
    write_output(kept, removed, args.out)

    flagged = sum(1 for r in kept if r["issues"])
    print(f"Read {len(rows)} rows from {args.input.name}")
    print(f"  kept:    {len(kept)} ({flagged} flagged for review)")
    print(f"  removed: {len(removed)}")
    for r in removed:
        print(f"    row {r['source_row']:>3}: {r['reason']}")
    print(f"Saved {args.out.name}")


if __name__ == "__main__":
    main()
