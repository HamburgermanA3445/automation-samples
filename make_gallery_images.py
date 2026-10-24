"""Build the gig gallery PNGs from the real sample outputs.

Run the three samples first, then:
    python make_gallery_images.py

Images are 2560x1538 (2x Fiverr's 1280x769) and go into each sample's
screenshots/ folder.
"""

import csv
from pathlib import Path

import openpyxl
import pypdfium2 as pdfium
from matplotlib import font_manager
from PIL import Image, ImageDraw, ImageFilter, ImageFont

ROOT = Path(__file__).parent
CANVAS = (2560, 1538)
BACKGROUND = "#e6e9ed"
SCALE = 2  # draw spreadsheet panels at 2x for sharp text


# --- Fonts --------------------------------------------------------------------

def font(size, bold=False):
    candidates = ["calibrib.ttf", "arialbd.ttf"] if bold else ["calibri.ttf", "arial.ttf"]
    for name in candidates:
        path = Path("C:/Windows/Fonts") / name
        if path.exists():
            return ImageFont.truetype(str(path), size)
    fallback = font_manager.findfont(f"DejaVu Sans:{'bold' if bold else 'normal'}")
    return ImageFont.truetype(fallback, size)


# --- Spreadsheet-style table ----------------------------------------------------

def col_letter(i):
    return chr(ord("A") + i)


def fit_text(draw, text, fnt, max_w):
    """Cut text to fit the cell, the way a spreadsheet clips it."""
    while text and draw.textlength(text, font=fnt) > max_w:
        text = text[:-1]
    return text


def spreadsheet_panel(columns, rows, filename, tabs, active_tab, fills=None,
                      styled_header=True, blank_rows=1):
    """Draw a table that looks like a sheet open in Excel.

    columns: list of (header, width_px, align) with align "l" or "r".
    rows: list of lists of display strings.
    fills: {(row_index, col_index): "#hex"} for highlighted cells.
    """
    s = SCALE
    fills = fills or {}
    rownum_w, colhdr_h, row_h, title_h, tabs_h = 44 * s, 22 * s, 25 * s, 34 * s, 30 * s
    width = rownum_w + sum(w * s for _, w, _ in columns)
    n_rows = len(rows) + 1 + blank_rows
    height = title_h + colhdr_h + row_h * n_rows + tabs_h

    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    regular, bold, small = font(15 * s), font(15 * s, bold=True), font(13 * s)
    grid, chrome, chrome_line, chrome_text = "#dadcdf", "#f3f4f6", "#c9ccd1", "#5f6368"

    # File name bar
    d.rectangle([0, 0, width, title_h], fill="#fafafa")
    d.line([0, title_h - 1, width, title_h - 1], fill=chrome_line, width=s)
    d.text((14 * s, title_h // 2), filename, font=font(14 * s, bold=True), fill="#3c4043", anchor="lm")

    # Column letters
    y0 = title_h
    d.rectangle([0, y0, width, y0 + colhdr_h], fill=chrome)
    x = rownum_w
    for i, (_, w, _) in enumerate(columns):
        w *= s
        d.text((x + w // 2, y0 + colhdr_h // 2), col_letter(i), font=small, fill=chrome_text, anchor="mm")
        d.line([x, y0, x, y0 + colhdr_h], fill=chrome_line, width=1)
        x += w
    d.line([0, y0 + colhdr_h - 1, width, y0 + colhdr_h - 1], fill=chrome_line, width=s)

    # Cells
    body_top = y0 + colhdr_h
    all_rows = [[h for h, _, _ in columns]] + rows + [[""] * len(columns)] * blank_rows
    for r, values in enumerate(all_rows):
        top = body_top + r * row_h
        d.rectangle([0, top, rownum_w - 1, top + row_h], fill=chrome)
        d.text((rownum_w // 2, top + row_h // 2), str(r + 1), font=small, fill=chrome_text, anchor="mm")
        x = rownum_w
        for c, ((_, w, align), value) in enumerate(zip(columns, values)):
            w *= s
            is_header = r == 0
            fill = fills.get((r - 1, c))
            if is_header and styled_header:
                fill = "#dde3ea"
            if fill:
                d.rectangle([x, top, x + w, top + row_h], fill=fill)
            fnt = bold if is_header and styled_header else regular
            text = fit_text(d, str(value), fnt, w - 12 * s)
            if align == "r" and not is_header:
                d.text((x + w - 6 * s, top + row_h // 2), text, font=fnt, fill="#202124", anchor="rm")
            else:
                d.text((x + 6 * s, top + row_h // 2), text, font=fnt, fill="#202124", anchor="lm")
            x += w
        d.line([0, top + row_h, width, top + row_h], fill=grid, width=1)

    x = rownum_w
    body_bottom = body_top + row_h * n_rows
    for _, w, _ in columns:
        d.line([x, body_top, x, body_bottom], fill=grid, width=1)
        x += w * s
    d.line([rownum_w - 1, y0, rownum_w - 1, body_bottom], fill=chrome_line, width=s)

    # Sheet tabs
    d.rectangle([0, body_bottom, width, height], fill=chrome)
    d.line([0, body_bottom, width, body_bottom], fill=chrome_line, width=s)
    x = rownum_w + 10 * s
    for tab in tabs:
        fnt = bold if tab == active_tab else regular
        tw = int(d.textlength(tab, font=fnt)) + 28 * s
        if tab == active_tab:
            d.rectangle([x, body_bottom, x + tw, height - 4 * s], fill="white")
            d.line([x + 8 * s, height - 6 * s, x + tw - 8 * s, height - 6 * s], fill="#1e7145", width=3 * s)
        d.text((x + tw // 2, body_bottom + tabs_h // 2 - s), tab, font=fnt, fill="#202124", anchor="mm")
        x += tw + 4 * s

    d.rectangle([0, 0, width - 1, height - 1], outline="#b9bdc3", width=s)
    return img


# --- PDF pages and layout -------------------------------------------------------

def render_pdf_page(path, height):
    pdf = pdfium.PdfDocument(str(path))
    page = pdf[0]
    scale = height / page.get_height()
    img = page.render(scale=scale).to_pil().convert("RGB")
    pdf.close()
    return img


def paste_with_shadow(canvas, img, xy, radius=18, offset=(0, 8)):
    shadow = Image.new("L", (img.width + radius * 4, img.height + radius * 4), 0)
    ImageDraw.Draw(shadow).rectangle(
        [radius * 2, radius * 2, radius * 2 + img.width, radius * 2 + img.height], fill=70)
    shadow = shadow.filter(ImageFilter.GaussianBlur(radius))
    sx, sy = xy[0] - radius * 2 + offset[0], xy[1] - radius * 2 + offset[1]
    canvas.paste(Image.new("RGB", shadow.size, "#5a6270"), (sx, sy), shadow)
    canvas.paste(img, xy)


def fit(img, max_w, max_h):
    ratio = min(max_w / img.width, max_h / img.height, 1.0)
    if ratio < 1:
        img = img.resize((int(img.width * ratio), int(img.height * ratio)), Image.LANCZOS)
    return img


def save_centered(panel_or_panels, path, margin=70):
    """Fit one image (or a vertical stack) onto the gallery canvas."""
    canvas = Image.new("RGB", CANVAS, BACKGROUND)
    panels = panel_or_panels if isinstance(panel_or_panels, list) else [panel_or_panels]
    ratio = min((CANVAS[0] - 2 * margin) / max(p.width for p in panels),
                (CANVAS[1] - 2 * margin - 40 * (len(panels) - 1)) / sum(p.height for p in panels), 1.0)
    panels = [p.resize((int(p.width * ratio), int(p.height * ratio)), Image.LANCZOS) if ratio < 1 else p
              for p in panels]
    total_h = sum(p.height for p in panels) + 40 * (len(panels) - 1)
    x = (CANVAS[0] - max(p.width for p in panels)) // 2  # stacked panels share a left edge
    y = (CANVAS[1] - total_h) // 2
    for p in panels:
        paste_with_shadow(canvas, p, (x, y))
        y += p.height + 40
    path.parent.mkdir(exist_ok=True)
    canvas.save(path, optimize=True)
    print(f"saved {path.relative_to(ROOT)}")


def side_by_side(images, path, margin=70, gap=60):
    canvas = Image.new("RGB", CANVAS, BACKGROUND)
    max_h = CANVAS[1] - 2 * margin
    total_w = sum(i.width * max_h / i.height for i in images) + gap * (len(images) - 1)
    ratio = min(1.0, (CANVAS[0] - 2 * margin) / total_w)
    images = [i.resize((int(i.width * max_h / i.height * ratio), int(max_h * ratio)), Image.LANCZOS)
              for i in images]
    x = (CANVAS[0] - sum(i.width for i in images) - gap * (len(images) - 1)) // 2
    for img in images:
        paste_with_shadow(canvas, img, (x, (CANVAS[1] - img.height) // 2))
        x += img.width + gap
    canvas.save(path, optimize=True)
    print(f"saved {path.relative_to(ROOT)}")


def fanned_pages(pages, path, margin=70):
    """Overlapping pages, left to right, so different layouts are visible."""
    canvas = Image.new("RGB", CANVAS, BACKGROUND)
    h = CANVAS[1] - 2 * margin
    pages = [p.resize((int(p.width * h / p.height), h), Image.LANCZOS) for p in pages]
    step = (CANVAS[0] - 2 * margin - pages[0].width) // (len(pages) - 1)
    for i, page in enumerate(pages):
        paste_with_shadow(canvas, page, (margin + i * step, margin), radius=22, offset=(-6, 6))
    canvas.save(path, optimize=True)
    print(f"saved {path.relative_to(ROOT)}")


# --- Helpers for reading the outputs back -----------------------------------------

def fmt(value, kind=None):
    if value is None:
        return ""
    if kind == "money":
        return f"{value:,.2f}"
    if kind == "date":
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# --- The three samples ------------------------------------------------------------

def pdf_invoices_gallery():
    base = ROOT / "pdf-invoices-to-excel"
    shots = base / "screenshots"
    shots.mkdir(exist_ok=True)

    names = ["Northfield_NP-10517.pdf", "BrightLine_Invoice_0883.pdf", "greenway inv 2026-117.pdf"]
    pages = [render_pdf_page(base / "sample_invoices" / n, 1500) for n in names]
    fanned_pages(pages, shots / "01_sample_invoices.png")

    wb = openpyxl.load_workbook(base / "output" / "invoices.xlsx")
    tabs = wb.sheetnames
    ws = wb["Invoices"]
    rows = [[r[0], r[1], r[2], fmt(r[3], "date"), fmt(r[4], "money"), fmt(r[5], "money"),
             fmt(r[6], "money"), fmt(r[7]), r[8]] for r in ws.iter_rows(min_row=2, values_only=True)]
    invoices = spreadsheet_panel(
        [("File", 205, "l"), ("Vendor", 205, "l"), ("Invoice No", 105, "l"),
         ("Invoice Date", 100, "l"), ("Subtotal", 80, "r"), ("Tax", 65, "r"),
         ("Total", 80, "r"), ("Lines", 50, "r"), ("Check", 60, "l")],
        rows, "invoices.xlsx", tabs, "Invoices")

    ws = wb["Line Items"]
    rows = [[r[0], r[1], fmt(r[2], "date"), r[3], fmt(r[4]), fmt(r[5], "money"), fmt(r[6], "money")]
            for r in ws.iter_rows(min_row=2, max_row=11, values_only=True)]
    items = spreadsheet_panel(
        [("Vendor", 205, "l"), ("Invoice No", 105, "l"), ("Invoice Date", 100, "l"),
         ("Description", 310, "l"), ("Qty", 45, "r"), ("Unit Price", 85, "r"), ("Amount", 85, "r")],
        rows, "invoices.xlsx", tabs, "Line Items", blank_rows=0)
    save_centered([invoices, items], shots / "02_excel_output.png")


def branded_invoices_gallery():
    base = ROOT / "csv-to-branded-invoices"
    shots = base / "screenshots"
    shots.mkdir(exist_ok=True)

    out = base / "output"
    first = render_pdf_page(out / "HSB-1042_copper-kettle-cafe.pdf", 1500)
    second = render_pdf_page(out / "HSB-1043_juniper-lane-market.pdf", 1500)
    side_by_side([first, second], shots / "01_invoice_pdfs.png")

    with open(base / "orders.csv", newline="", encoding="utf-8") as f:
        data = list(csv.reader(f))
    sheet = spreadsheet_panel(
        [(data[0][0], 88, "l"), (data[0][1], 160, "l"), (data[0][2], 120, "l"),
         (data[0][3], 100, "l"), (data[0][4], 165, "l"), (data[0][5], 40, "r"), (data[0][6], 75, "r")],
        data[1:], "orders.csv", ["orders"], "orders", styled_header=False)
    side_by_side([sheet, first], shots / "02_orders_to_invoice.png", gap=90)


def cleanup_gallery():
    base = ROOT / "sheet-cleanup-dedupe"
    shots = base / "screenshots"
    shots.mkdir(exist_ok=True)

    with open(base / "messy_contacts.csv", newline="", encoding="utf-8") as f:
        data = list(csv.reader(f))
    widths = [150, 175, 285, 135, 110, 55, 100]
    before = spreadsheet_panel([(h, w, "l") for h, w in zip(data[0], widths)], data[1:],
                               "messy_contacts.csv", ["messy_contacts"], "messy_contacts",
                               styled_header=False, blank_rows=0)
    save_centered(before, shots / "01_before.png")

    wb = openpyxl.load_workbook(base / "output" / "contacts_clean.xlsx")
    ws = wb["Contacts"]
    rows, fills = [], {}
    for i, row in enumerate(ws.iter_rows(min_row=2)):
        values = [c.value for c in row]
        rows.append([fmt(values[0]), fmt(values[1]), fmt(values[2]), fmt(values[3]), fmt(values[4]),
                     fmt(values[5]), fmt(values[6], "date"), fmt(values[7]), fmt(values[8])])
        for j, cell in enumerate(row):
            if cell.fill.fgColor.rgb == "00FFF2B3":
                fills[(i, j)] = "#fff2b3"
    after = spreadsheet_panel(
        [("Name", 135, "l"), ("Company", 170, "l"), ("Email", 290, "l"), ("Phone", 140, "l"),
         ("City", 105, "l"), ("State", 45, "l"), ("Last Contact", 95, "l"), ("Source Row", 90, "r"),
         ("Issues", 125, "l")],
        rows, "contacts_clean.xlsx", wb.sheetnames, "Contacts", fills=fills, blank_rows=0)

    ws = wb["Removed Rows"]
    removed_rows = [[fmt(v) for v in r] for r in ws.iter_rows(min_row=2, values_only=True)]
    removed = spreadsheet_panel(
        [("Source Row", 90, "r"), ("Name (as entered)", 140, "l"), ("Email (as entered)", 230, "l"),
         ("Reason", 600, "l"), ("Kept Row", 80, "r")],
        removed_rows, "contacts_clean.xlsx", wb.sheetnames, "Removed Rows", blank_rows=0)
    save_centered([after, removed], shots / "02_after.png", margin=50)


if __name__ == "__main__":
    pdf_invoices_gallery()
    branded_invoices_gallery()
    cleanup_gallery()
